#!/usr/bin/env python3
"""
Project Butler - Office File Converter (Enhanced)
Converts LLM-generated Markdown/CSV into professional .pptx / .docx / .xlsx files.

Usage:
  python convert.py <input.md/csv> <output.ext> <type>
  python convert.py ppt <input.json> <output.pptx>  (legacy, auto-detected)

Types: ppt, word, excel
Input can be: raw Markdown (.md), CSV (.csv), or JSON with {"content":"...","topic":"..."}
"""

import json, sys, csv, re, io, os
from pathlib import Path

# ─── PPTX ──────────────────────────────────────────────────────────
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PPT_C = {
    "bg":     RGBColor(0x0D, 0x11, 0x1A),
    "light":  RGBColor(0x15, 0x1A, 0x28),
    "accent": RGBColor(0x00, 0xD4, 0xAA),
    "accent2":RGBColor(0x6C, 0x8C, 0xFF),
    "white":  RGBColor(0xFF, 0xFF, 0xFF),
    "gray":   RGBColor(0x8A, 0x8F, 0x9A),
    "line":   RGBColor(0x2A, 0x30, 0x45),
}

# ─── DOCX ──────────────────────────────────────────────────────────
from docx import Document
from docx.shared import Inches as DI, Pt as DP, RGBColor as DRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOC_C = {
    "title":   DRGB(0x0D, 0x11, 0x1A),
    "heading": DRGB(0x15, 0x1A, 0x28),
    "accent":  DRGB(0x00, 0x8A, 0x6E),
    "body":    DRGB(0x33, 0x33, 0x33),
    "light":   DRGB(0x66, 0x66, 0x66),
    "line":    DRGB(0x00, 0xD4, 0xAA),
}

# ─── XLSX ──────────────────────────────────────────────────────────
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

XLSX_C = {
    "header_bg": "0D111A", "header_font": "FFFFFF",
    "even_row": "F0F4F8", "border": "D0D5DD", "accent": "00D4AA",
}

# ══════════════════════════════════════════════════════════════════════
#  PPTX GENERATION
# ══════════════════════════════════════════════════════════════════════

def create_pptx(content, topic, output_path):
    """Generate professional PPTX from Markdown content."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides_data = parse_ppt_markdown(content, topic)
    if not slides_data:
        slides_data = [{"title": topic or "Presentation", "points": [l for l in content.split("\n") if l.strip()]}]

    # Title Slide
    _add_ppt_slide_bg(prs, PPT_C["bg"])
    slide = prs.slides[-1]
    _add_accent_line(slide, 0, 0, 13.333, 0.06, PPT_C["accent"])
    _add_textbox(slide, 1.2, 2.0, 10.9, 2.0, topic or "Project Report", 44, PPT_C["white"], True)
    _add_textbox(slide, 1.2, 4.2, 10.9, 1.2, "Project Butler — AI-Powered Document Generation", 20, PPT_C["gray"])
    _add_accent_line(slide, 1.2, 6.8, 2.5, 0.04, PPT_C["accent2"])

    # Content Slides
    for sd in slides_data:
        _add_ppt_slide_bg(prs, PPT_C["bg"])
        slide = prs.slides[-1]
        _add_accent_line(slide, 0, 0, 13.333, 0.06, PPT_C["accent"])

        title = re.sub(r'^(?:第\s*\d+\s*页[：:]\s*|Page\s*\d+[：:]\s*)', '', sd.get("title", ""))
        if title:
            _add_textbox(slide, 1.0, 0.4, 11.3, 1.0, title, 32, PPT_C["white"], True)
        _add_accent_line(slide, 1.0, 1.35, 2.0, 0.04, PPT_C["accent"])

        points = sd.get("points", [])
        if points:
            txBox = slide.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.3), Inches(5.0))
            tf = txBox.text_frame; tf.word_wrap = True
            for i, pt in enumerate(points):
                pt = re.sub(r'^[-•]\s*', '', pt.strip())
                if not pt or pt.startswith("**视觉建议**") or pt.startswith("👁"):
                    continue
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"  ▸ {pt}"
                p.font.size = Pt(18); p.font.color.rgb = PPT_C["white"]
                p.font.name = "Microsoft YaHei"; p.space_after = Pt(14)

        num = sd.get("num", 0)
        _add_textbox(slide, 11.5, 7.0, 1.5, 0.4, f"{num} / {len(slides_data)}", 10, PPT_C["gray"], align=PP_ALIGN.RIGHT)

    prs.save(output_path)

def _add_ppt_slide_bg(prs, color):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color
    return slide

def _add_accent_line(slide, l, t, w, h, color):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()

def _add_textbox(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txBox.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    p.font.name = "Microsoft YaHei"; p.alignment = align

def parse_ppt_markdown(content, topic):
    """Parse Markdown into structured slides using ## headers as slide breaks."""
    slides = []
    blocks = re.split(r'\n(?=##\s)', content) if '## ' in content else [content]

    for block in blocks:
        block = block.strip()
        if not block: continue
        lines = block.split("\n")
        slide = {"title": "", "points": [], "num": len(slides) + 1}

        for line in lines:
            line = line.strip()
            if not line: continue
            if line.startswith("## "):
                slide["title"] = line[3:].strip()
            elif line.startswith("# "):
                slide["title"] = line[2:].strip()
            elif line.startswith("- ") or line.startswith("* ") or line.startswith("• "):
                slide["points"].append(line)
            elif line and not line.startswith("```") and not line.startswith("|"):
                slide["points"].append(line)

        if slide["title"] or slide["points"]:
            slides.append(slide)

    return slides

# ══════════════════════════════════════════════════════════════════════
#  DOCX GENERATION
# ══════════════════════════════════════════════════════════════════════

def create_docx(content, topic, output_path):
    """Generate professional DOCX from Markdown content."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"; style.font.size = DP(12); style.font.color.rgb = DOC_C["body"]

    for s in doc.sections:
        s.top_margin = DI(1.0); s.bottom_margin = DI(1.0)
        s.left_margin = DI(1.2); s.right_margin = DI(1.2)

    # Cover
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(topic or "Project Document"); r.font.size = DP(28); r.font.bold = True
    r.font.color.rgb = DOC_C["title"]

    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("━" * 40); r2.font.size = DP(10); r2.font.color.rgb = DOC_C["line"]

    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Generated by Project Butler AI Agent"); r3.font.size = DP(11)
    r3.font.color.rgb = DOC_C["light"]; r3.font.italic = True
    doc.add_paragraph()

    # Content parsing
    in_code = False; code_lines = []
    for line in content.split("\n"):
        s = line.strip()
        if s.startswith("```"):
            if in_code:
                if code_lines:
                    p = doc.add_paragraph()
                    for cl in code_lines:
                        run = p.add_run(cl + "\n")
                        run.font.name = "Consolas"; run.font.size = DP(10)
                code_lines = []; in_code = False
            else: in_code = True
            continue
        if in_code: code_lines.append(line); continue
        if not s: continue

        if s.startswith("### "):
            p = doc.add_paragraph(); r = p.add_run(s[4:])
            r.font.size = DP(16); r.font.bold = True; r.font.color.rgb = DOC_C["heading"]
        elif s.startswith("## "):
            p = doc.add_paragraph(); r = p.add_run(s[3:])
            r.font.size = DP(20); r.font.bold = True; r.font.color.rgb = DOC_C["accent"]
            p.space_before = DP(18)
        elif s.startswith("# "):
            p = doc.add_paragraph(); r = p.add_run(s[2:])
            r.font.size = DP(24); r.font.bold = True; r.font.color.rgb = DOC_C["title"]
        elif s.startswith("- ") or s.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _render_md_run(p, s[2:])
        elif s.startswith("**"):
            p = doc.add_paragraph(); p.space_before = DP(12)
            text = s.replace("**", "")
            if "：" in text or ": " in text:
                parts = re.split(r'[：:]\s*', text, maxsplit=1)
                r = p.add_run(parts[0] + "："); r.font.bold = True
                r.font.size = DP(13); r.font.color.rgb = DOC_C["heading"]
                if len(parts) > 1:
                    r2 = p.add_run(parts[1]); r2.font.size = DP(12)
            else:
                r = p.add_run(text); r.font.bold = True; r.font.size = DP(13)
        else:
            p = doc.add_paragraph(); _render_md_run(p, s)

    doc.save(output_path)

def _render_md_run(p, text):
    """Render inline bold/italic/code in a paragraph."""
    pattern = r'(\*\*.*?\*\*|\*.*?\*|`.*?`)'
    parts = re.split(pattern, text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.font.bold = True
        elif part.startswith("*") and part.endswith("*"):
            r = p.add_run(part[1:-1]); r.font.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = DP(10)
        else:
            p.add_run(part)

# ══════════════════════════════════════════════════════════════════════
#  XLSX GENERATION
# ══════════════════════════════════════════════════════════════════════

def create_xlsx(content, topic, output_path):
    """Generate professional XLSX from CSV/Markdown table content."""
    wb = Workbook(); ws = wb.active; ws.title = "Financial Model"

    # Try parsing CSV first, then Markdown table
    rows = list(csv.reader(io.StringIO(content)))
    if len(rows) <= 1:
        # Try Markdown table parsing
        rows = _parse_md_table(content)
    if not rows:
        ws["A1"] = "No parseable data"; wb.save(output_path); return

    # Styles
    h_font = Font(name="Microsoft YaHei", size=11, bold=True, color=XLSX_C["header_font"])
    h_fill = PatternFill(start_color=XLSX_C["header_bg"], end_color=XLSX_C["header_bg"], fill_type="solid")
    h_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    d_font = Font(name="Microsoft YaHei", size=10)
    e_fill = PatternFill(start_color=XLSX_C["even_row"], end_color=XLSX_C["even_row"], fill_type="solid")
    thin = Side(style="thin", color=XLSX_C["border"])
    medium = Side(style="medium", color=XLSX_C["accent"])
    t_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    a_border = Border(left=thin, right=thin, top=thin, bottom=medium)

    for ci, h in enumerate(rows[0], 1):
        cell = ws.cell(row=1, column=ci, value=h.strip().strip('"'))
        cell.font = h_font; cell.fill = h_fill; cell.alignment = h_align; cell.border = a_border

    for ri, row in enumerate(rows[1:], 2):
        for ci, val in enumerate(row, 1):
            v = val.strip().strip('"')
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.font = d_font; cell.alignment = Alignment(vertical="center"); cell.border = t_border
            if ri % 2 == 0: cell.fill = e_fill
            try:
                float(v)
                cell.number_format = '#,##0.0' if '.' in v else '#,##0'
            except ValueError: pass

    # Column widths
    for ci in range(1, min(len(rows[0]) + 1, 20)):
        max_w = max((sum(2 if ord(c) > 127 else 1 for c in str(ws.cell(row=ri, column=ci).value or ""))
                     for ri in range(1, min(len(rows) + 1, 51))), default=10)
        ws.column_dimensions[get_column_letter(ci)].width = min(max_w + 4, 45)

    ws.freeze_panes = "A2"
    try:
        ws.auto_filter.ref = f"A1:{get_column_letter(min(len(rows[0]), 20))}{len(rows)}"
    except: pass

    wb.save(output_path)

def _parse_md_table(content):
    """Parse Markdown table into list of rows."""
    rows = []
    in_table = False
    for line in content.split("\n"):
        s = line.strip()
        if s.startswith("|") and s.endswith("|"):
            if not in_table:
                in_table = True
            if re.match(r'^[\|\s\-:]+$', s):
                continue
            cells = [c.strip() for c in s.split("|")[1:-1]]
            rows.append(cells)
        elif in_table and not s.startswith("|"):
            in_table = False
    return rows

# ══════════════════════════════════════════════════════════════════════

def main():
    args = sys.argv[1:]

    # Auto-detect mode: <input> <output> <type> vs <type> <input> <output>
    if len(args) >= 3:
        if args[0] in ("ppt", "docx", "xlsx", "word", "excel"):
            # Legacy mode: convert.py <type> <input> <output>
            tool = args[0]; input_path = args[1]; output_path = args[2]
        else:
            # New mode: convert.py <input> <output> <type>
            input_path = args[0]; output_path = args[1]; tool = args[2]
    else:
        print(json.dumps({"error": "Usage: convert.py <input> <output> <ppt|word|excel>"}))
        sys.exit(1)

    # Read input
    ext = Path(input_path).suffix.lower()
    with open(input_path, "r", encoding="utf-8") as f:
        raw = f.read()

    # Parse content
    if ext == ".json":
        data = json.loads(raw)
        content = data.get("content", raw)
        topic = data.get("topic", "Project Document")
    else:
        content = raw
        # Extract first heading as topic
        m = re.search(r'^#\s+(.+)$', content, re.MULTILINE)
        topic = m.group(1).strip() if m else Path(input_path).stem

    # Ensure output dir
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    # Dispatch
    try:
        tool_lower = tool.lower()
        if tool_lower in ("ppt", "pptx"):
            create_pptx(content, topic, output_path)
        elif tool_lower in ("word", "docx"):
            create_docx(content, topic, output_path)
        elif tool_lower in ("excel", "xlsx"):
            create_xlsx(content, topic, output_path)
        else:
            print(json.dumps({"error": f"Unknown type: {tool}"})); sys.exit(1)

        print(json.dumps({
            "ok": True, "file": os.path.abspath(output_path),
            "size_bytes": os.path.getsize(output_path), "type": tool_lower
        }))
    except Exception as e:
        print(json.dumps({"error": str(e)})); sys.exit(1)

if __name__ == "__main__":
    main()
